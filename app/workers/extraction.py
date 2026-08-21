"""Extraction + RAG воркер (Ф6) — ОТДЕЛЬНЫЙ systemd-сервис.

Не живёт в uvicorn-процессе осознанно (adversarial-ревью плана): парсинг
PDF/DOCX — CPU-bound и заморозил бы API (workers=1). Здесь парсинг уходит
в thread pool, цикл — каждые 30 секунд:

1. text_extraction_jobs (pending) → извлечение текста файла (pypdf /
   python-docx / plain) → search_documents.body_text — «поиск по
   содержимому файлов» начинает находить документы.
2. RAG-reconcile (если AI настроен): search_documents → rag_chunks
   с эмбеддингами; без ключа шаг тихо пропускается.

Скан очередей — bypass; доменная запись — в tenant-scoped сессиях.
Run via systemd: `signaris-hub[-staging]-extraction.service` (Restart=always).
"""

from __future__ import annotations

import asyncio
from datetime import UTC, datetime
from pathlib import Path

import structlog
from sqlalchemy import select, text

from app import log as log_config
from app.db import tenant_scoped_session
from app.models.library import LibraryMaterial, MaterialVersion
from app.models.search_document import TextExtractionJob
from app.services.library_storage import absolute_path
from app.services.llm import LLMEmbeddingsUnsupported, LLMNotConfigured, get_provider
from app.services.rag_indexer import reconcile
from app.services.search_indexer import upsert_document

log = structlog.get_logger("workers.extraction")

POLL_SEC = 30
BATCH = 10
MAX_ATTEMPTS = 3
TEXT_LIMIT = 200_000

_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
# xlsx: листы читаем потоково (read_only), ограничения — чтобы воркер с
# MemoryMax=512M не распух на «простыне» в сотни тысяч строк.
XLSX_MAX_ROWS = 5_000
XLSX_MAX_BYTES = 20 * 1024 * 1024
EXTRACTABLE_MIMES = frozenset({"application/pdf", _DOCX_MIME, _XLSX_MIME})


def _extract_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    # Таблицы — построчно, ячейки через табуляцию (python-docx не включает их
    # в paragraphs; техкарты и чек-листы живут именно там).
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append("\t".join(cells))
    return "\n".join(parts)


def _extract_xlsx(path: Path) -> str:
    if path.stat().st_size > XLSX_MAX_BYTES:
        raise ValueError("xlsx больше 20 МБ — текст не извлекаем")
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    parts: list[str] = []
    rows_left = XLSX_MAX_ROWS
    try:
        for ws in wb.worksheets:
            parts.append(f"## {ws.title}")
            for row in ws.iter_rows(values_only=True):
                if rows_left <= 0:
                    break
                cells = ["" if v is None else str(v).strip() for v in row]
                if any(cells):
                    parts.append("\t".join(cells).rstrip("\t"))
                    rows_left -= 1
            if rows_left <= 0:
                break
    finally:
        wb.close()
    return "\n".join(parts)


def _extract_text_sync(path: Path, mime: str) -> str:
    """CPU-bound парсинг — зовётся через asyncio.to_thread."""
    if mime == "application/pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        pages = [page.extract_text() or "" for page in reader.pages[:200]]
        return "\n".join(pages)[:TEXT_LIMIT]
    if mime == _DOCX_MIME:
        return _extract_docx(path)[:TEXT_LIMIT]
    if mime == _XLSX_MIME:
        return _extract_xlsx(path)[:TEXT_LIMIT]
    if mime.startswith("text/"):
        return path.read_text(errors="ignore")[:TEXT_LIMIT]
    raise ValueError(f"Извлечение текста для {mime} не поддерживается")


async def _process_extraction_batch() -> int:
    async with tenant_scoped_session(None, bypass_rls=True) as scan:
        jobs = (
            (
                await scan.execute(
                    select(TextExtractionJob)
                    .where(
                        TextExtractionJob.status == "pending",
                        TextExtractionJob.attempts < MAX_ATTEMPTS,
                    )
                    .order_by(TextExtractionJob.created_at)
                    .limit(BATCH)
                )
            )
            .scalars()
            .all()
        )
        job_specs = [
            (j.id, j.tenant_id, j.object_id, j.storage_key, j.mime) for j in jobs
        ]

    done = 0
    for job_id, tenant_id, object_id, storage_key, mime in job_specs:
        async with tenant_scoped_session(tenant_id) as session:
            job = await session.get(TextExtractionJob, job_id)
            if job is None or job.status != "pending":
                continue
            try:
                path = absolute_path(storage_key)
                body_text = await asyncio.to_thread(_extract_text_sync, path, mime)
                # Текст живёт в ВЕРСИИ независимо от статуса материала: индекс
                # берёт его оттуда при каждой переиндексации (раньше PATCH
                # материала затирал body_text, а черновик не индексировался
                # никогда — ОС 2026-08, 86/181 материалов прода без текста).
                version = (
                    await session.execute(
                        select(MaterialVersion).where(
                            MaterialVersion.storage_key == storage_key
                        )
                    )
                ).scalar_one_or_none()
                if version is not None:
                    version.extracted_text = body_text or None
                    version.extracted_at = datetime.now(UTC)
                material = await session.get(LibraryMaterial, object_id)
                if material is not None and material.status == "published":
                    await upsert_document(
                        session,
                        tenant_id=material.tenant_id,
                        object_type="library_material",
                        object_id=material.id,
                        title=material.title,
                        snippet=material.description,
                        body_text=body_text,
                        audience_id=material.audience_id,
                        published_at=material.published_at,
                        url_path=f"/learn/library?m={material.id}",
                    )
                job.status = "done"
                job.error = None
                done += 1
            except Exception as e:  # noqa: BLE001 — джоба не должна ронять цикл
                job.attempts += 1
                job.error = str(e)[:500]
                if job.attempts >= MAX_ATTEMPTS:
                    job.status = "failed"
                log.warning(
                    "extraction.job_failed",
                    job_id=job_id,
                    attempts=job.attempts,
                    err=str(e)[:200],
                )
            await session.commit()
    return done


async def _enqueue_backfill() -> int:
    """Стартовый шаг: поставить джобы версиям, которые воркер ещё не извлекал
    (`extracted_at IS NULL`) и у которых нет pending-джобы — бэкфилл после 0041
    и для форматов, которые раньше падали (xlsx). Старые failed/done джобы не
    трогаем: новая встаёт в очередь с теми же storage_key/mime."""
    async with tenant_scoped_session(None, bypass_rls=True) as session:
        # extracted_at штампуется ВСЕГДА (и при пустом тексте) — иначе сканы
        # без текста вставали бы в очередь при каждом старте воркера.
        rows = (
            (
                await session.execute(
                    select(MaterialVersion)
                    .where(
                        MaterialVersion.extracted_at.is_(None),
                        MaterialVersion.mime.in_(EXTRACTABLE_MIMES),
                        ~select(TextExtractionJob.id)
                        .where(
                            TextExtractionJob.storage_key == MaterialVersion.storage_key,
                            TextExtractionJob.status == "pending",
                        )
                        .exists(),
                    )
                    .order_by(MaterialVersion.created_at)
                    .limit(300)
                )
            )
            .scalars()
            .all()
        )
        for v in rows:
            session.add(
                TextExtractionJob(
                    tenant_id=v.tenant_id,
                    object_type="library_material",
                    object_id=v.material_id,
                    storage_key=v.storage_key,
                    mime=v.mime,
                )
            )
        await session.commit()
        return len(rows)


async def _process_rag() -> None:
    try:
        provider = get_provider()
    except LLMNotConfigured:
        return  # AI не настроен — extraction работает, RAG ждёт ключа
    async with tenant_scoped_session(None, bypass_rls=True) as scan:
        tenant_ids = [
            r[0]
            for r in await scan.execute(
                text("SELECT DISTINCT tenant_id FROM search_documents")
            )
        ]
    for tenant_id in tenant_ids:
        async with tenant_scoped_session(tenant_id) as session:
            try:
                stats = await reconcile(session, provider, limit=5)
            except LLMEmbeddingsUnsupported:
                # Провайдер без embeddings (DeepSeek): ассистент работает
                # через лексический retrieval, вект-индекс не строим.
                return
            await session.commit()
            if stats["docs"] or stats["orphans_deleted"]:
                log.info("rag.reconciled", tenant_id=str(tenant_id), **stats)


async def main() -> None:
    log_config.configure()
    log.info("extraction.worker_started", poll_sec=POLL_SEC)
    try:
        queued = await _enqueue_backfill()
        if queued:
            log.info("extraction.backfill_enqueued", jobs=queued)
    except Exception as e:  # noqa: BLE001 — бэкфилл не должен ронять старт
        log.error("extraction.backfill_failed", err=str(e)[:300])
    while True:
        try:
            processed = await _process_extraction_batch()
            if processed:
                log.info("extraction.batch_done", jobs=processed)
            await _process_rag()
        except Exception as e:  # noqa: BLE001 — цикл живёт всегда
            log.error("extraction.cycle_failed", err=str(e)[:300])
        await asyncio.sleep(POLL_SEC)


if __name__ == "__main__":
    asyncio.run(main())
