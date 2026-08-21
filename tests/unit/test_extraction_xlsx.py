"""Извлечение текста из xlsx/docx для поиска и предпросмотра (ОС 2026-08)."""

from __future__ import annotations

from pathlib import Path

from app.workers.extraction import _XLSX_MIME, _extract_text_sync


def test_xlsx_extracts_all_sheets_as_tab_rows(tmp_path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Техкарта"
    ws.append(["Блюдо", "Выход, г", "Цена"])
    ws.append(["Азу из говядины", 250, 390])
    ws.append([None, None, None])
    ws2 = wb.create_sheet("Чек-лист")
    ws2.append(["Открытие", "Да"])
    path = tmp_path / "t.xlsx"
    wb.save(path)

    text = _extract_text_sync(path, _XLSX_MIME)
    assert "## Техкарта" in text and "## Чек-лист" in text
    assert "Азу из говядины\t250\t390" in text
    assert "Открытие\tДа" in text
    # Пустые строки листа не попадают в текст.
    assert all(line.strip() for line in text.splitlines())


def test_docx_includes_tables(tmp_path: Path) -> None:
    import docx

    d = docx.Document()
    d.add_paragraph("Стандарт сервиса")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Шаг"
    table.rows[0].cells[1].text = "Приветствие"
    path = tmp_path / "t.docx"
    d.save(path)
    text = _extract_text_sync(
        path, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert "Стандарт сервиса" in text
    assert "Шаг\tПриветствие" in text
